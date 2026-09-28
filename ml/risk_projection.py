from __future__ import annotations

from pathlib import Path
from typing import Any, Mapping, cast

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .config import TARGET_LABELS


# =============================================================================
# IMPORTANTE
# =============================================================================
#
# Estos valores son índices matemáticos definidos para el proyecto.
# NO representan coeficientes médicos ni probabilidades clínicas validadas.
# =============================================================================

MISSING_RISK_PENALTY = 0.10
MAX_ADJUSTED_RISK = 0.99

HORIZONS = {
    "24h": 1,
    "7d": 7,
    "30d": 30,
}

CSV_ENCODING = "utf-8-sig"


# =============================================================================
# Estimador
# =============================================================================

def _get_model(pipeline: Any) -> Any:
    """
    Obtiene el estimador final de un sklearn Pipeline.

    La arquitectura actual utiliza el step `estimator`.
    Se conserva compatibilidad con `model` para artefactos antiguos.
    """
    if hasattr(pipeline, "named_steps"):
        named_steps = pipeline.named_steps

        if "estimator" in named_steps:
            return named_steps["estimator"]

        if "model" in named_steps:
            return named_steps["model"]

    return pipeline


def _get_classes(
    pipeline: Any,
    fallback: Any = None,
) -> list[Any]:
    """
    Obtiene las clases del estimador final.
    """
    model = _get_model(pipeline)

    classes = getattr(
        model,
        "classes_",
        None,
    )

    if classes is None:
        classes = getattr(
            pipeline,
            "classes_",
            None,
        )

    if classes is None:
        classes = fallback

    if classes is None:
        return []

    try:
        return list(classes)
    except Exception:
        return []


# =============================================================================
# Conversión de probabilidades a índice de riesgo
# =============================================================================

def _positive_class_index(
    classes: list[Any],
) -> int | None:
    """
    Determina la clase positiva para clasificación binaria.

    Convención técnica:
    1. Si las clases son numéricas, se utiliza la mayor.
    2. Si son strings, se intenta localizar SI/TRUE/YES/Y.
    3. Como fallback, se utiliza la última clase.

    Esta convención NO implica significado clínico.
    """
    if len(classes) != 2:
        return None

    # -------------------------------------------------------------------------
    # Clases numéricas
    # -------------------------------------------------------------------------

    try:
        numeric = pd.to_numeric(
            pd.Series(classes),
            errors="coerce",
        ).to_numpy(dtype=float)

        if np.isfinite(numeric).all():
            return int(
                np.argmax(numeric)
            )
    except Exception:
        pass

    # -------------------------------------------------------------------------
    # Convención binaria textual
    # -------------------------------------------------------------------------

    positive_labels = {
        "SI",
        "SÍ",
        "TRUE",
        "VERDADERO",
        "YES",
        "Y",
        "1",
        "1.0",
    }

    normalized = [
        str(value).strip().upper()
        for value in classes
    ]

    for index, value in enumerate(normalized):
        if value in positive_labels:
            return index

    # -------------------------------------------------------------------------
    # Fallback
    # -------------------------------------------------------------------------

    return 1


def _risk_from_probabilities(
    probabilities: Any,
    classes: Any,
    target_key: str | None = None,
) -> np.ndarray:
    """
    Convierte probabilidades de clasificación en un índice matemático [0, 1].

    Binario:
        utiliza la probabilidad de la clase positiva.

    Multiclase ordinal:
        calcula el promedio ponderado de las probabilidades según el orden
        numérico de las clases.

    Multiclase no ordinal:
        utiliza la probabilidad máxima.

    El resultado NO es una probabilidad clínica validada.
    """
    del target_key

    probability = np.asarray(
        probabilities,
        dtype=float,
    )

    class_values = list(classes)

    if probability.ndim != 2:
        return np.zeros(
            probability.shape[0]
            if probability.ndim > 0
            else 0,
            dtype=float,
        )

    if probability.shape[1] == 0:
        return np.zeros(
            probability.shape[0],
            dtype=float,
        )

    if len(class_values) != probability.shape[1]:
        class_values = list(
            range(
                probability.shape[1]
            )
        )

    # -------------------------------------------------------------------------
    # Binario
    # -------------------------------------------------------------------------

    if len(class_values) == 2:
        positive_index = _positive_class_index(
            class_values
        )

        if positive_index is None:
            positive_index = 1

        positive_index = min(
            max(
                positive_index,
                0,
            ),
            probability.shape[1] - 1,
        )

        return np.clip(
            probability[:, positive_index],
            0.0,
            1.0,
        )

    # -------------------------------------------------------------------------
    # Multiclase ordinal
    # -------------------------------------------------------------------------

    try:
        numeric_classes = pd.to_numeric(
            pd.Series(class_values),
            errors="coerce",
        ).to_numpy(dtype=float)

        if (
            len(numeric_classes)
            == probability.shape[1]
            and np.isfinite(
                numeric_classes
            ).all()
            and numeric_classes.max()
            > numeric_classes.min()
        ):
            weights = (
                numeric_classes
                - numeric_classes.min()
            ) / (
                numeric_classes.max()
                - numeric_classes.min()
            )

            risk = probability @ weights

            return np.clip(
                risk,
                0.0,
                1.0,
            )

    except Exception:
        pass

    # -------------------------------------------------------------------------
    # Fallback
    # -------------------------------------------------------------------------

    return np.clip(
        np.max(
            probability,
            axis=1,
        ),
        0.0,
        1.0,
    )


# =============================================================================
# Datos faltantes
# =============================================================================

def _missing_fraction(
    X: pd.DataFrame,
) -> np.ndarray:
    """
    Calcula la fracción de variables faltantes por paciente.
    """
    if X.empty:
        return np.zeros(
            len(X),
            dtype=float,
        )

    return (
        X.isna()
        .mean(axis=1)
        .to_numpy(dtype=float)
    )


def _adjust_risk_for_missing_data(
    base_risk: np.ndarray,
    missing_fraction: np.ndarray,
) -> np.ndarray:
    """
    Ajusta el índice de riesgo por ausencia de datos.

    La penalización representa incertidumbre definida para el proyecto.
    No es un coeficiente clínico.
    """
    base = np.asarray(
        base_risk,
        dtype=float,
    )

    missing = np.asarray(
        missing_fraction,
        dtype=float,
    )

    if len(base) != len(missing):
        raise ValueError(
            "base_risk y missing_fraction deben "
            "tener la misma longitud."
        )

    adjusted = (
        base
        + missing * MISSING_RISK_PENALTY
    )

    return np.clip(
        adjusted,
        0.0,
        MAX_ADJUSTED_RISK,
    )


# =============================================================================
# Proyección temporal
# =============================================================================

def _project_risk(
    adjusted_risk: float,
    days: int,
) -> float:
    """
    Proyecta el índice mediante hazard diario constante.

    R(t) = 1 - (1 - R24)^t

    Supuestos matemáticos:
    - hazard diario constante;
    - independencia entre días;
    - el índice inicial se interpreta como riesgo base diario.

    Esta proyección NO es una estimación clínica validada.
    """
    if days <= 0:
        return 0.0

    risk = float(
        np.clip(
            adjusted_risk,
            0.0,
            MAX_ADJUSTED_RISK,
        )
    )

    projected = (
        1.0
        - (1.0 - risk) ** days
    )

    return float(
        np.clip(
            projected,
            0.0,
            1.0,
        )
    )


# =============================================================================
# Identificación de pacientes
# =============================================================================

def _patient_ids(
    df: pd.DataFrame,
) -> pd.Series:
    """
    Obtiene identificadores de pacientes preservando el índice original.
    """
    candidate_columns = [
        "ID",
        "patient_id",
        "paciente_id",
    ]

    patient_column = next(
        (
            column
            for column in candidate_columns
            if column in df.columns
        ),
        None,
    )

    fallback = pd.Series(
        np.arange(
            1,
            len(df) + 1,
        ),
        index=df.index,
    )

    if patient_column is not None:
        ids = df[
            patient_column
        ].copy()

        return ids.where(
            ids.notna(),
            fallback,
        )

    return fallback


# =============================================================================
# Inferencia
# =============================================================================

def _safe_predict_proba(
    pipeline: Any,
    X: pd.DataFrame,
) -> tuple[np.ndarray, list[Any]] | None:
    """
    Ejecuta predict_proba y devuelve probabilidades y clases alineadas.
    """
    if not hasattr(
        pipeline,
        "predict_proba",
    ):
        return None

    try:
        probabilities = np.asarray(
            pipeline.predict_proba(X),
            dtype=float,
        )
    except Exception:
        return None

    if (
        probabilities.ndim != 2
        or probabilities.shape[0] != len(X)
    ):
        return None

    classes = _get_classes(
        pipeline
    )

    if not classes:
        if probabilities.shape[1] == 2:
            classes = [0, 1]
        else:
            classes = list(
                range(
                    probabilities.shape[1]
                )
            )

    if len(classes) != probabilities.shape[1]:
        return None

    return probabilities, classes


# =============================================================================
# Cálculo por target
# =============================================================================

def _append_target_rows(
    rows: list[dict[str, Any]],
    df: pd.DataFrame,
    result: Any,
    key: str,
    features: list[str],
    patient_ids: pd.Series,
) -> None:
    """
    Calcula y agrega las proyecciones de un target.
    """
    pipeline = getattr(
        result,
        "_risk_pipeline",
        None,
    )

    if pipeline is None:
        return

    if not features:
        return

    missing_features = [
        feature
        for feature in features
        if feature not in df.columns
    ]

    if missing_features:
        return

    X = df[
        features
    ].copy()

    prediction_data = _safe_predict_proba(
        pipeline,
        X,
    )

    if prediction_data is None:
        return

    probabilities, classes = prediction_data

    base_risk = _risk_from_probabilities(
        probabilities,
        classes,
        target_key=key,
    )

    missing = _missing_fraction(
        X
    )

    adjusted = _adjust_risk_for_missing_data(
        base_risk,
        missing,
    )

    completeness = np.clip(
        1.0 - missing,
        0.0,
        1.0,
    )

    if not (
        len(base_risk)
        == len(missing)
        == len(adjusted)
        == len(completeness)
        == len(patient_ids)
    ):
        return

    label = TARGET_LABELS.get(
        key,
        getattr(
            result,
            "column",
            key,
        ),
    )

    model_name = getattr(
        result,
        "selected",
        None,
    )

    for index, patient_id in enumerate(
        patient_ids
    ):
        adjusted_value = float(
            adjusted[index]
        )

        row: dict[str, Any] = {
            "ID_paciente": patient_id,
            "target": key,
            "targetLabel": label,
            "model": model_name,
            "riesgo_modelo_pct": round(
                float(
                    base_risk[index] * 100.0
                ),
                4,
            ),
            "datos_faltantes_pct": round(
                float(
                    missing[index] * 100.0
                ),
                4,
            ),
            "completitud_datos_pct": round(
                float(
                    completeness[index]
                    * 100.0
                ),
                4,
            ),
            "riesgo_ajustado_incertidumbre_pct": round(
                adjusted_value * 100.0,
                4,
            ),
        }

        for horizon, days in HORIZONS.items():
            projected = _project_risk(
                adjusted_value,
                days,
            )

            row[
                f"riesgo_{horizon}"
            ] = round(
                projected * 100.0,
                4,
            )

        rows.append(row)


# =============================================================================
# Documentación de fórmulas
# =============================================================================

def _formula_rows(
    target_results: Mapping[str, Any],
) -> list[dict[str, Any]]:
    """
    Construye documentación de las fórmulas utilizadas.
    """
    rows: list[dict[str, Any]] = []

    for key, result in target_results.items():
        if not getattr(
            result,
            "artifact",
            None,
        ):
            continue

        rows.append(
            {
                "target": key,
                "targetLabel": TARGET_LABELS.get(
                    key,
                    getattr(
                        result,
                        "column",
                        key,
                    ),
                ),
                "model": getattr(
                    result,
                    "selected",
                    None,
                ),
                "penalizacion_datos_faltantes": (
                    MISSING_RISK_PENALTY
                ),
                "limite_riesgo_ajustado": (
                    MAX_ADJUSTED_RISK
                ),
                "formula_riesgo_ajustado": (
                    "min(0.99, "
                    "riesgo_modelo + "
                    "0.10 * "
                    "fraccion_datos_faltantes)"
                ),
                "formula_24h": (
                    "R24 = riesgo_ajustado"
                ),
                "formula_7d": (
                    "R7 = 1 - "
                    "(1 - riesgo_ajustado)^7"
                ),
                "formula_30d": (
                    "R30 = 1 - "
                    "(1 - riesgo_ajustado)^30"
                ),
                "supuesto_proyeccion": (
                    "Hazard diario constante "
                    "e independencia entre días."
                ),
                "interpretacion": (
                    "Índice matemático de proyección. "
                    "No constituye una fórmula médica "
                    "ni una probabilidad clínica validada."
                ),
            }
        )

    return rows


# =============================================================================
# Excel: riesgos
# =============================================================================

def _write_risk_sheet(
    workbook: Workbook,
    df_risk: pd.DataFrame,
) -> None:
    """
    Escribe y formatea la hoja principal de riesgos.
    """
    # `active` está tipado como Worksheet | None por los stubs de openpyxl.
    # En un Workbook recién creado existe una hoja activa.
    worksheet = cast(
        Worksheet,
        workbook.active,
    )

    worksheet.title = "Riesgos"

    headers = list(
        df_risk.columns
    )

    worksheet.append(headers)

    for row in df_risk.itertuples(
        index=False,
        name=None,
    ):
        worksheet.append(
            list(row)
        )

    worksheet.freeze_panes = "A2"

    if worksheet.max_row >= 2:
        worksheet.auto_filter.ref = (
            worksheet.dimensions
        )

    for cell in worksheet[1]:
        cell.font = Font(
            bold=True
        )
        cell.alignment = Alignment(
            wrap_text=True,
            horizontal="center",
            vertical="center",
        )

    worksheet.row_dimensions[
        1
    ].height = 35

    for column_index, _column in enumerate(
        worksheet.columns,
        start=1,
    ):
        values = [
            cell.value
            for cell in worksheet[
                get_column_letter(column_index)
            ][:100]
        ]

        width = max(
            [
                len(
                    str(
                        value or ""
                    )
                )
                for value in values
            ]
            + [12]
        )

        worksheet.column_dimensions[
            get_column_letter(
                column_index
            )
        ].width = min(
            width + 2,
            38,
        )


# =============================================================================
# Excel: fórmulas
# =============================================================================

def _write_formula_sheet(
    workbook: Workbook,
    formula_rows: list[dict[str, Any]],
) -> None:
    """
    Escribe documentación de fórmulas y supuestos.
    """
    worksheet: Worksheet = workbook.create_sheet(
        "Formulas"
    )

    if not formula_rows:
        worksheet.append(
            ["Información"]
        )
        worksheet.append(
            ["No hay proyecciones disponibles."]
        )
        return

    headers = list(
        formula_rows[0].keys()
    )

    worksheet.append(headers)

    for row in formula_rows:
        worksheet.append(
            [
                row.get(header)
                for header in headers
            ]
        )

    for cell in worksheet[1]:
        cell.font = Font(
            bold=True
        )
        cell.alignment = Alignment(
            wrap_text=True,
            horizontal="center",
        )

    worksheet.freeze_panes = "A2"

    for index, header in enumerate(
        headers,
        start=1,
    ):
        width = 20

        if (
            "formula" in header.lower()
            or "interpretacion" in header.lower()
            or "supuesto" in header.lower()
        ):
            width = 65

        worksheet.column_dimensions[
            get_column_letter(index)
        ].width = width


# =============================================================================
# Excel: curvas
# =============================================================================

def _write_risk_charts(
    workbook: Workbook,
    df_risk: pd.DataFrame,
) -> None:
    """
    Genera una hoja con datos y curvas de riesgo proyectado por target.
    """
    worksheet: Worksheet = workbook.create_sheet(
        "Curvas_Riesgo"
    )

    if df_risk.empty:
        return

    row_cursor = 1

    for target, group in df_risk.groupby(
        "target",
        sort=False,
    ):
        target_text = str(target)

        worksheet.cell(
            row_cursor,
            1,
            f"Target: {target_text}",
        )

        headers = [
            "ID_paciente",
            "24h",
            "7d",
            "30d",
        ]

        for column_index, header in enumerate(
            headers,
            start=1,
        ):
            worksheet.cell(
                row_cursor + 1,
                column_index,
                header,
            )

        for row_index, (_, row) in enumerate(
            group.iterrows(),
            start=row_cursor + 2,
        ):
            worksheet.cell(
                row_index,
                1,
                row["ID_paciente"],
            )
            worksheet.cell(
                row_index,
                2,
                row["riesgo_24h"],
            )
            worksheet.cell(
                row_index,
                3,
                row["riesgo_7d"],
            )
            worksheet.cell(
                row_index,
                4,
                row["riesgo_30d"],
            )

        for cell in worksheet[
            row_cursor + 1
        ]:
            cell.font = Font(
                bold=True
            )

        data_start = row_cursor + 1
        data_end = (
            row_cursor
            + 1
            + len(group)
        )

        if data_end >= data_start + 1:
            chart = LineChart()

            chart.title = (
                "Curva de riesgo proyectado — "
                f"{target_text}"
            )
            chart.y_axis.title = "%"
            chart.x_axis.title = "Paciente"

            chart.add_data(
                Reference(
                    worksheet,
                    min_col=2,
                    max_col=4,
                    min_row=data_start,
                    max_row=data_end,
                ),
                titles_from_data=True,
            )

            chart.set_categories(
                Reference(
                    worksheet,
                    min_col=1,
                    min_row=data_start + 1,
                    max_row=data_end,
                )
            )

            chart.height = 8
            chart.width = 18

            worksheet.add_chart(
                chart,
                f"F{row_cursor + 1}",
            )

        row_cursor += (
            len(group)
            + 28
        )

    worksheet.column_dimensions[
        "A"
    ].width = 20

    for column in [
        "B",
        "C",
        "D",
    ]:
        worksheet.column_dimensions[
            column
        ].width = 14


# =============================================================================
# Exportación
# =============================================================================

def export_risk_projections(
    output_dir: Path,
    df: pd.DataFrame,
    target_results: Mapping[str, Any],
    feature_map: Mapping[str, list[str]],
) -> Path | None:
    """
    Genera proyecciones matemáticas de riesgo por paciente.

    El pipeline entrenado se reutiliza directamente para mantener las mismas
    transformaciones utilizadas durante entrenamiento.

    Retorna la ruta del XLSX generado o None si no hay datos disponibles.
    """
    if df is None or df.empty:
        return None

    output = (
        Path(output_dir)
        / "resultado_riesgos"
    )

    output.mkdir(
        parents=True,
        exist_ok=True,
    )

    rows: list[dict[str, Any]] = []

    patient_ids = _patient_ids(
        df
    )

    for key, result in target_results.items():
        if not getattr(
            result,
            "artifact",
            None,
        ):
            continue

        features = feature_map.get(
            key
        )

        if not features:
            continue

        try:
            _append_target_rows(
                rows=rows,
                df=df,
                result=result,
                key=key,
                features=list(features),
                patient_ids=patient_ids,
            )
        except Exception:
            # Un target defectuoso no debe impedir generar los restantes.
            continue

    risk_frame = pd.DataFrame(
        rows
    )

    if risk_frame.empty:
        return None

    formula_rows = _formula_rows(
        target_results
    )

    xlsx_path = (
        output
        / "riesgos_por_paciente.xlsx"
    )

    workbook = Workbook()

    _write_risk_sheet(
        workbook,
        risk_frame,
    )

    _write_formula_sheet(
        workbook,
        formula_rows,
    )

    _write_risk_charts(
        workbook,
        risk_frame,
    )

    workbook.save(
        xlsx_path
    )

    # -------------------------------------------------------------------------
    # CSV principal
    # -------------------------------------------------------------------------

    risk_frame.to_csv(
        output
        / "riesgos_por_paciente.csv",
        index=False,
        encoding=CSV_ENCODING,
    )

    # -------------------------------------------------------------------------
    # CSV de fórmulas
    # -------------------------------------------------------------------------

    pd.DataFrame(
        formula_rows
    ).to_csv(
        output
        / "formulas_riesgo.csv",
        index=False,
        encoding=CSV_ENCODING,
    )

    return xlsx_path
